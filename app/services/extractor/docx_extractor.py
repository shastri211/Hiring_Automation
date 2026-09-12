import docx
from .base import BaseExtractor

class DOCXExtractor(BaseExtractor):
    async def extract(self, file_path: str) -> str:
        try:
            doc = docx.Document(file_path)
            text_content = []
            
            # Extract paragraphs
            for para in doc.paragraphs:
                text_content.append(para.text)
                
            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_data = []
                    for cell in row.cells:
                        row_data.append(cell.text.replace("\n", " ").strip())
                    text_content.append(" | ".join(row_data))
                    
            return "\n".join(text_content).strip()
        except Exception as e:
            raise Exception(f"DOCX extraction failed: {str(e)}")
